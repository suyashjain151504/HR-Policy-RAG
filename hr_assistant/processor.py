"""Parse raw binary files in S3 to plain-text JSON, once.

    raw/<prefix>/file.pdf|docx|pptx  --parse-->  processed/<prefix>/file.json

Called by ingest so PDFs/DOCX/PPTX are parsed a single time, not on every
vector-store rebuild. document_loader.py then reads the JSON back.
.txt needs no special parser beyond UTF-8 decode.
"""

import io
import json
import logging

import boto3
from botocore.config import Config as BotoConfig
from docx import Document as DocxReader
from pptx import Presentation
from pypdf import PdfReader

from hr_assistant import config
from hr_assistant.document_loader import extract_policy_category

logger = logging.getLogger(__name__)

_RAW_TO_PROCESSED = {
    config.S3_PREFIX: config.PROCESSED_HR_PREFIX,
    config.NOISE_S3_PREFIX: config.PROCESSED_NOISE_PREFIX,
}


def _s3_client():
    return boto3.client(
        "s3",
        endpoint_url=config.AWS_ENDPOINT_URL,
        region_name=config.AWS_DEFAULT_REGION,
        aws_access_key_id=config.AWS_ACCESS_KEY_ID,
        aws_secret_access_key=config.AWS_SECRET_ACCESS_KEY,
        config=BotoConfig(s3={"addressing_style": "path"}),
    )


def _parse_pdf(raw_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(raw_bytes: bytes) -> str:
    doc = DocxReader(io.BytesIO(raw_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def _parse_pptx(raw_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(raw_bytes))
    lines = []
    for slide in prs.slides:
        if slide.shapes.title is not None:
            lines.append(slide.shapes.title.text)
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


_PARSERS = {".pdf": _parse_pdf, ".docx": _parse_docx, ".pptx": _parse_pptx}


def parse_object(key: str, raw_bytes: bytes) -> str:
    """Turn object bytes into plain text, based on the file extension."""
    ext = "." + key.rsplit(".", 1)[-1].lower() if "." in key else ""
    if ext == ".txt":
        return raw_bytes.decode("utf-8")
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"No parser registered for file extension {ext!r} ({key})")
    return parser(raw_bytes)


def process_raw_to_json(bucket_name: str = config.S3_BUCKET_NAME) -> int:
    """Parse every raw object and write a processed JSON record to S3.

    Returns how many records were written. Each record is what
    load_processed_documents_from_s3() reads back.
    """
    client = _s3_client()
    count = 0
    for raw_prefix, processed_prefix in _RAW_TO_PROCESSED.items():
        token = None
        while True:
            kwargs = {"Bucket": bucket_name, "Prefix": raw_prefix}
            if token:
                kwargs["ContinuationToken"] = token
            page = client.list_objects_v2(**kwargs)
            for obj in page.get("Contents") or []:
                key = obj["Key"]
                filename = key.rsplit("/", 1)[-1]
                if "." not in filename:
                    continue
                raw_bytes = client.get_object(Bucket=bucket_name, Key=key)["Body"].read()
                text = parse_object(key, raw_bytes)
                record = {
                    "source": filename,
                    "policy_category": extract_policy_category(text),
                    "text": text,
                    "raw_s3_path": f"s3://{bucket_name}/{key}",
                }
                processed_name = filename.rsplit(".", 1)[0] + ".json"
                processed_key = f"{processed_prefix}{processed_name}"
                client.put_object(
                    Bucket=bucket_name,
                    Key=processed_key,
                    Body=json.dumps(record, indent=2).encode("utf-8"),
                    ContentType="application/json",
                )
                logger.info("  %s  ->  %s", key, processed_key)
                count += 1
            if not page.get("IsTruncated"):
                break
            token = page.get("NextContinuationToken")
    return count
